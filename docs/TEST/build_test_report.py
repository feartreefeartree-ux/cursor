#!/usr/bin/env python3
"""Convert 04_FAN_F系统测试报告.docx into a plain Excel workbook."""

import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# Column character widths for wrap-height estimate (approx Excel units).
COL_W = {1: 14, 2: 28, 3: 18, 4: 58, 5: 8}

SCOPE = [
    ("类别", "数量", "结果"),
    ("CAN/ISO-TP", 5, "全部通过"),
    ("UDS", 18, "全部通过"),
    ("安全访问", 5, "全部通过"),
    ("系统检查", 1, "全部通过"),
    ("OTA", 11, "全部通过"),
]

ARTIFACTS = [
    ("制品",),
    ("FAN_F_ECU_TEST.hex",),
    ("FAN_CAN_APP_SAFE_RAMP_UNIVERSAL.hex",),
    ("rfd1_nxp_sdk_ref_NONPROD_TEST_ONLY.pkg",),
    ("FAN_F_CAN_Tester.exe",),
]

CAN_LOG = [
    ("项目", "结果"),
    ("总帧数", 12290),
    ("TX", 11789),
    ("RX", 501),
    ("CAN ID数量", 11),
    ("CAN周期", "count=5, min=483ms, max=541ms, expected=500±100ms"),
]

DID = [
    ("用例", "DID", "ASCII", "响应", "结果"),
    (
        "UDS-018",
        "F187双通道精确值读取",
        "A61300000859 + 0x20",
        "62 F1 87 41 36 31 33 30 30 30 30 30 38 35 39 20",
        "通过",
    ),
    (
        "UDS-019",
        "F189双通道精确值读取",
        "A6130A14A8S001001",
        "62 F1 89 41 36 31 33 30 41 31 34 41 38 53 30 30 31 30 30 31",
        "通过",
    ),
    (
        "UDS-020",
        "F18A双通道精确值读取",
        "A1448",
        "62 F1 8A 41 31 34 34 38",
        "通过",
    ),
    (
        "UDS-021",
        "F191双通道精确值读取",
        "A6130A1448H001001",
        "62 F1 91 41 36 31 33 30 41 31 34 34 38 48 30 30 31 30 30 31",
        "通过",
    ),
    (
        "UDS-022",
        "F197双通道精确值读取",
        "FAN_F + 9个空格",
        "62 F1 97 46 41 4E 5F 46 20 20 20 20 20 20 20 20 20",
        "通过",
    ),
]

DID_NOTE = (
    "上述5项均在OBD、OTA物理通道完成逐字节比较。"
    "F187末尾0x20和F197尾部空格计入长度及字节校验。"
    "F18F仍按未就绪状态处理，读取返回NRC 0x31。"
)

# Word split this across two tables; keep one header and all cases.
FULL = [
    ("编号", "测试项目", "类别", "主要实测结果", "结论"),
    ("CAN-001", "CAN设备连接", "CAN/ISO-TP", "CAN Broker已连接", "通过"),
    (
        "CAN-002",
        "OBD/OTA双通道响应",
        "CAN/ISO-TP",
        "62 F1 97 46 41 4E 5F 46 20 20 20 20 20 20 20 20 20",
        "通过",
    ),
    ("CAN-003", "ISO-TP单帧", "CAN/ISO-TP", "Boot诊断服务已在线", "通过"),
    (
        "CAN-004",
        "ISO-TP多帧",
        "CAN/ISO-TP",
        "62 F1 97 46 41 4E 5F 46 20 20 20 20 20 20 20 20 20",
        "通过",
    ),
    ("UDS-001", "Default会话", "UDS", "Boot诊断服务已在线", "通过"),
    ("UDS-002", "Programming会话门禁", "UDS", "50 02 00 32 01 F4", "通过"),
    ("UDS-003", "Extended会话和P2/P2*", "UDS", "50 03 00 32 01 F4", "通过"),
    ("UDS-004", "S3超时回Default", "UDS", "S3超时后Boot已回到Default会话：7F 27 7E", "通过"),
    ("UDS-005", "TesterPresent和抑制正响应", "UDS", "在200ms观察窗内无响应", "通过"),
    (
        "UDS-007",
        "系统标识DID读取",
        "UDS",
        "62 F1 97 46 41 4E 5F 46 20 20 20 20 20 20 20 20 20",
        "通过",
    ),
    (
        "UDS-018",
        "F187双通道精确值读取",
        "UDS",
        "62 F1 87 41 36 31 33 30 30 30 30 30 38 35 39 20",
        "通过",
    ),
    (
        "UDS-019",
        "F189双通道精确值读取",
        "UDS",
        "62 F1 89 41 36 31 33 30 41 31 34 41 38 53 30 30 31 30 30 31",
        "通过",
    ),
    ("UDS-020", "F18A双通道精确值读取", "UDS", "62 F1 8A 41 31 34 34 38", "通过"),
    (
        "UDS-021",
        "F191双通道精确值读取",
        "UDS",
        "62 F1 91 41 36 31 33 30 41 31 34 34 38 48 30 30 31 30 30 31",
        "通过",
    ),
    (
        "UDS-022",
        "F197双通道精确值读取",
        "UDS",
        "62 F1 97 46 41 4E 5F 46 20 20 20 20 20 20 20 20 20",
        "通过",
    ),
    ("UDS-008", "DID写入门禁", "UDS", "7F 2E 7F", "通过"),
    ("UDS-009", "DTC读取", "UDS", "59 02 09", "通过"),
    ("UDS-012", "不支持SID，NRC11", "UDS", "7F 99 11", "通过"),
    ("UDS-013", "不支持子功能，NRC12", "UDS", "7F 10 12", "通过"),
    ("UDS-014", "长度错误，NRC13", "UDS", "7F 10 13", "通过"),
    ("UDS-015", "请求顺序错误，NRC24", "UDS", "7F 36 24", "通过"),
    (
        "SEC-001",
        "Level 1正确Key",
        "安全访问",
        "Level 1解锁；Provider=FAN_F_Seed2Key_TEST_x64；SHA256=0e3596b4eb8f3...",
        "通过",
    ),
    (
        "SEC-002",
        "Level 3正确Key",
        "安全访问",
        "Level 3解锁；Provider=FAN_F_Seed2Key_TEST_x64；SHA256=0e3596b4eb8f3...",
        "通过",
    ),
    ("SEC-003", "错误Key和NRC35", "安全访问", "7F 27 **REDACTED**", "通过"),
    ("SEC-004", "无Seed直接发Key", "安全访问", "7F 27 **REDACTED**", "通过"),
    ("SEC-008", "会话降级清除权限", "安全访问", "7F 14 33", "通过"),
    (
        "SYS-010",
        "最终诊断与系统标识检查",
        "系统检查",
        "62 F1 97 46 41 4E 5F 46 20 20 20 20 20 20 20 20 20",
        "通过",
    ),
    (
        "UDS-006",
        "ECUReset响应、复位和重连",
        "UDS",
        "ECU在第8次探测首次恢复；最终连续在线1500ms，复核4次；稳定窗口重启0次，复位/切换期间探测无响应7次（最终已恢复）",
        "通过",
    ),
    ("OTA-001", "非编程会话RequestDownload", "OTA", "7F 34 7F", "通过"),
    (
        "OTA-002",
        "F002正常条件",
        "OTA",
        "OTA-002通过：正常边界：响应=71 01 F0 02 00，判定=允许；请求前条件快照已刷新；vehicle_speed...",
        "通过",
    ),
    (
        "OTA-003",
        "F002速度禁止",
        "OTA",
        "OTA-003通过：车速超限：响应=71 01 F0 02 01，判定=禁止；请求前条件快照已刷新；vehicle_speed...",
        "通过",
    ),
    (
        "OTA-004",
        "F002电源禁止",
        "OTA",
        "OTA-004通过：电源模式不满足：响应=71 01 F0 02 01，判定=禁止；请求前条件快照已刷新；vehicle_sp...",
        "通过",
    ),
    (
        "OTA-005",
        "F002驻车/发动机禁止",
        "OTA",
        "OTA-005通过：驻车条件不满足：响应=71 01 F0 02 01，判定=禁止；请求前条件快照已刷新；vehicle_sp...",
        "通过",
    ),
    (
        "OTA-006",
        "F002缺失/无效/超时策略",
        "OTA",
        "OTA-006通过：未接收：响应=71 01 F0 02 00，判定=允许；请求前条件快照已刷新；engine_speed帧龄...",
        "通过",
    ),
    (
        "OTA-007",
        "Flash Driver正常下载与认证",
        "OTA",
        "OTA-007 Flash Driver下载认证及F184 READY门禁通过，未发送FF00；Boot诊断F197在线",
        "通过",
    ),
    (
        "OTA-010",
        "合法RequestDownload与地址长度门禁",
        "OTA",
        "OTA-010合法RequestDownload通过，错误地址、零长度和越界长度均被拒绝",
        "通过",
    ),
    (
        "OTA-009",
        "FF00擦除受控目标",
        "OTA",
        "OTA-009 FF00完成；未下载、未提交候选；复位后原可用运行态恢复：Boot诊断F197在线",
        "通过",
    ),
    (
        "OTA-028",
        "正向完整在线升级",
        "OTA",
        "ECU受控非活动区刷写、应用运行验证与Trial确认完成",
        "通过",
    ),
    (
        "CAN-005",
        "CAN周期",
        "CAN/ISO-TP",
        "count=5, min=483ms, max=541ms, expected=500±100ms",
        "通过",
    ),
    ("OTA-040", "OBD/OTA安全隔离", "OTA", "7F 37 7F", "通过"),
]

THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def font(size=10, bold=False):
    return Font(name="微软雅黑", size=size, bold=bold, color="000000")


def align(h="left", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def display_width(value):
    n = 0
    for ch in str(value):
        n += 2 if ord(ch) > 127 else 1
    return n


def height_for(values, col_widths, min_h=20, line_h=15, max_h=64):
    lines = 1
    for value, width in zip(values, col_widths):
        if value is None or value == "":
            continue
        usable = max(4, width - 1.2)
        lines = max(lines, math.ceil(display_width(value) / (usable * 1.15)))
    return min(max_h, max(min_h, 8 + lines * line_h))


def put(ws, row, col, value, *, bold=False, size=10, h="left", border=False):
    cell = ws.cell(row, col, value)
    cell.font = font(size=size, bold=bold)
    cell.alignment = align(h)
    if border:
        cell.border = THIN
    if isinstance(value, int):
        cell.alignment = align("center")
        cell.number_format = "#,##0" if value >= 1000 else "0"
    return cell


def write_table(ws, start, rows, col_count, last_span=None):
    """last_span: merge the last value across extra columns so long text stays one line."""
    widths = [COL_W[c] for c in range(1, col_count + 1)]
    if last_span and last_span > col_count:
        widths[-1] = sum(COL_W[c] for c in range(col_count, last_span + 1))
    for i, row in enumerate(rows):
        r = start + i
        header = i == 0
        for c, value in enumerate(row, 1):
            h = "center" if header or c == 1 or c == col_count else "left"
            if col_count <= 2 and c == 2 and not header:
                h = "left"
            if col_count == 3 and c in (2, 3) and not header:
                h = "center"
            if col_count == 1:
                h = "left"
            put(ws, r, c, value, bold=header, h=h, border=True)
        if last_span and last_span > col_count:
            ws.merge_cells(start_row=r, start_column=col_count, end_row=r, end_column=last_span)
            for c in range(col_count + 1, last_span + 1):
                ws.cell(r, c).border = THIN
                ws.cell(r, c).alignment = align("left" if not header else "center")
                ws.cell(r, c).font = font(10, header)
        ws.row_dimensions[r].height = height_for(row, widths, min_h=20 if header else 22)
    return start + len(rows)


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "测试报告"

    for col, w in COL_W.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.merge_cells("A1:E1")
    put(ws, 1, 1, "FAN_F ECU系统测试验证报告", bold=True, size=16, h="center")
    for c in range(2, 6):
        ws.cell(1, c).alignment = align("center")
        ws.cell(1, c).font = font(16, True)
    ws.row_dimensions[1].height = 28

    r = 3
    put(ws, r, 1, "1、 测试范围", bold=True, size=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 22
    r = write_table(ws, r + 1, SCOPE, 3) + 1

    put(ws, r, 1, "2 测试制品", bold=True, size=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 22
    r = write_table(ws, r + 1, ARTIFACTS, 1, last_span=4) + 1

    put(ws, r, 1, "3 CAN通信记录", bold=True, size=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 22
    r = write_table(ws, r + 1, CAN_LOG, 2, last_span=4) + 1

    put(ws, r, 1, "5 产品身份DID测试结果", bold=True, size=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 22
    r = write_table(ws, r + 1, DID, 5)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    put(ws, r, 1, DID_NOTE, size=10, h="left")
    ws.row_dimensions[r].height = height_for((DID_NOTE,), [sum(COL_W.values())], min_h=32, max_h=48)
    r += 2

    put(ws, r, 1, "6 全量测试结果", bold=True, size=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 22
    write_table(ws, r + 1, FULL, 5)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"

    out = "/workspace/docs/TEST/04_FAN_F系统测试报告.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    print("wrote", build())
